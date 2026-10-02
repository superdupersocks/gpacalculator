"""Download the federal source files for the admissions data into data/admissions/raw/.

    python3 scripts/admissions/fetch.py            # newest published year of each IPEDS file
    IPEDS_YEAR=2024 python3 scripts/admissions/fetch.py

Sources (Department of Education, public, no key needed):
- IPEDS complete data files (required): HD, ADM, IC, IC_AY, DRVEF, EF D, DRVGR and SFA, each from its newest
  year that has a dictionary (https://nces.ed.gov/ipeds/datacenter/data/<FILE>.zip and <FILE>_Dict.zip).
- IPEDS provisional release: NCES publishes a newer collection (e.g. 2024-25, with fall 2024 admissions) months
  before its complete data files, only through the Access database page and the data generator. A table listed
  in that release's Tablesdoc workbook replaces the complete file when it is newer; its variable titles and
  code labels come from the Tablesdoc. Set IPEDS_PROVISIONAL=0 to skip.
- College Scorecard, Most Recent Institution-Level Data (optional extras; https://collegescorecard.ed.gov/data/):
  SCORECARD_ZIP_FILE (a downloaded copy), else SCORECARD_ZIP_URL, else the release linked from the data page,
  else the pinned SCORECARD_ZIP.

Writes raw/scorecard/institutions.csv, raw/ipeds/<file>.csv, raw/ipeds/<file>_dict.json (variable titles and code
labels from the dictionary workbook) and data/admissions/manifest.json (URL, size, SHA-256, year of every file).
Needs openpyxl for the dictionary workbooks.
"""
import datetime
import hashlib
import io
import json
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile

sys.path.insert(0, os.path.dirname(__file__))
from common import OUT, RAW, write_json  # noqa: E402

# The Scorecard hosts refuse requests that don't look like a browser (403), so send a browser's headers.
HEADERS = {
    "User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "en-US,en;q=0.9",
    "Referer": "https://collegescorecard.ed.gov/data/",
}
IPEDS = "https://nces.ed.gov/ipeds/datacenter/data/"
ACCESS_PAGE = "https://nces.ed.gov/ipeds/use-the-data/download-access-database"
GENERATOR = "https://nces.ed.gov/ipeds/data-generator?year={year}&tableName={table}&HasRV=0&type=csv"
SCORECARD_PAGE = "https://collegescorecard.ed.gov/data/"
# "Most Recent Institution-Level Data", release of June 10, 2026. Update when Scorecard publishes a new release.
SCORECARD_ZIP = "https://ed-public-download.scorecard.network/downloads/Most-Recent-Cohorts-Institution_06102026.zip"
# key -> file name for a given year. ADM: admissions; HD: directory; IC: characteristics; IC_AY: tuition;
# DRVEF: derived enrollment; EF D: retention; DRVGR: derived graduation rates; SFA: net price.
IPEDS_FILES = {
    "hd": lambda y: f"HD{y}", "adm": lambda y: f"ADM{y}", "ic": lambda y: f"IC{y}", "ic_ay": lambda y: f"IC{y}_AY",
    "drvef": lambda y: f"DRVEF{y}", "efd": lambda y: f"EF{y}D", "drvgr": lambda y: f"DRVGR{y}",
    "sfa": lambda y: f"SFA{str(y - 1)[2:]}{str(y)[2:]}",
}


def get(url, tries=4, opener=None):
    for i in range(tries):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with (opener.open if opener else urllib.request.urlopen)(req, timeout=300) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            if e.code in (403, 404, 410):
                raise
            err = e
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            err = e
        time.sleep(2 ** (i + 1))
    raise err


def exists(url):
    try:
        req = urllib.request.Request(url, headers={**HEADERS, "Range": "bytes=0-1"})
        with urllib.request.urlopen(req, timeout=60) as r:
            return r.read(2) == b"PK"  # a zip, not an HTML "file not found" page
    except urllib.error.HTTPError:
        return False


def record(manifest, key, url, data, **extra):
    manifest["files"][key] = {"url": url, "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), **extra}


def pick_member(zf, suffix):
    """The file in a zip ending with suffix; IPEDS ships a revised '_rv' copy beside the provisional one."""
    names = [n for n in zf.namelist() if n.lower().endswith(suffix)]
    if not names:
        raise SystemExit(f"no *{suffix} in zip: {zf.namelist()}")
    revised = [n for n in names if "_rv" in n.lower()]
    return (revised or names)[0]


def parse_dictionary(xlsx_bytes):
    """Variable titles and code labels from an IPEDS dictionary workbook (sheets 'varlist' and 'Frequencies')."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    sheets = {ws.title.lower(): ws for ws in wb.worksheets}

    def rows(name):
        ws = sheets.get(name)
        if ws is None:
            return []
        it = ws.iter_rows(values_only=True)
        head = [str(h or "").strip().lower() for h in next(it)]
        return [dict(zip(head, r)) for r in it]

    out = {"vars": {}, "codes": {}}
    for r in rows("varlist"):
        if r.get("varname"):
            out["vars"][str(r["varname"]).upper()] = {"title": str(r.get("vartitle") or ""),
                                                      "type": str(r.get("datatype") or "")}
    for r in rows("frequencies"):
        if r.get("varname") is not None and r.get("codevalue") is not None:
            code = str(r["codevalue"]).strip()
            code = code[:-2] if code.endswith(".0") else code
            out["codes"].setdefault(str(r["varname"]).upper(), {})[code] = str(r.get("valuelabel") or "")
    if not out["vars"]:
        raise SystemExit(f"dictionary has no varlist sheet: {list(sheets)}")
    return out


def parse_tablesdoc(xlsx_bytes):
    """{TABLE: dictionary} from a release's Tablesdoc workbook (sheets vartable* and valuesets*), in the same
    shape as parse_dictionary."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)

    def rows(prefix):
        name = next((n for n in wb.sheetnames if n.lower().startswith(prefix)), None)
        if name is None:
            raise SystemExit(f"Tablesdoc has no {prefix} sheet: {wb.sheetnames}")
        it = wb[name].iter_rows(values_only=True)
        head = [str(h or "").strip().lower() for h in next(it)]
        return [dict(zip(head, r)) for r in it]

    out = {}
    for r in rows("vartable"):
        if r.get("tablename") and r.get("varname"):
            d = out.setdefault(str(r["tablename"]).upper(), {"vars": {}, "codes": {}})
            d["vars"][str(r["varname"]).upper()] = {"title": str(r.get("vartitle") or ""),
                                                    "type": str(r.get("datatype") or "")}
    for r in rows("valuesets"):
        if r.get("tablename") and r.get("varname") and r.get("codevalue") is not None:
            code = str(r["codevalue"]).strip()
            code = code[:-2] if code.endswith(".0") else code
            d = out.setdefault(str(r["tablename"]).upper(), {"vars": {}, "codes": {}})
            d["codes"].setdefault(str(r["varname"]).upper(), {})[code] = str(r.get("valuelabel") or "")
    return out


def newest_release(html):
    """(data year, Tablesdoc URL) of the newest collection on the Access database page; IPEDS202425 -> 2024."""
    links = re.findall(r"""href=["']([^"']*IPEDS(\d{4})(\d{2})Tablesdoc\.xlsx)["']""", html, re.I)
    if not links:
        return None, None
    url, start, _ = max(links, key=lambda t: t[1])
    return int(start), urllib.parse.urljoin(ACCESS_PAGE, url)


def provisional(manifest, years):
    """Tables from a release newer than the complete data files, via the data generator. Returns the keys
    replaced."""
    if os.environ.get("IPEDS_PROVISIONAL") == "0":
        return set()
    import http.cookiejar
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
    try:
        year, doc_url = newest_release(get(ACCESS_PAGE, tries=2, opener=opener).decode("utf-8", "ignore"))
        if year is None or all(year <= y for y in years.values()):
            print(f"IPEDS provisional: no release newer than the complete files ({year})")
            return set()
        doc = get(doc_url, opener=opener)
        tables = parse_tablesdoc(doc)
        # The data generator wants the session cookies a browser gets from the data files page.
        get(f"https://nces.ed.gov/ipeds/datacenter/DataFiles.aspx?year={year}&surveyNumber=1", tries=2,
            opener=opener)
    except Exception as e:
        print(f"WARNING: IPEDS provisional release not read ({e}); using complete data files only")
        manifest["ipeds_provisional_missing"] = str(e)
        return set()
    manifest["ipeds_provisional"] = {"year": year, "tablesdoc": doc_url}
    record(manifest, "ipeds_tablesdoc", doc_url, doc, year=year)
    done = set()
    for key, name_for in IPEDS_FILES.items():
        table = name_for(year).upper()
        if year <= years[key] or table not in tables:
            continue
        url = GENERATOR.format(year=year, table=table)
        try:
            data = get(url, opener=opener)
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                member = pick_member(zf, ".csv")
                csv_bytes = zf.read(member)
        except Exception as e:
            print(f"WARNING: IPEDS provisional {table} not downloaded ({e}); keeping {name_for(years[key])}")
            continue
        (RAW / "ipeds" / f"{key}.csv").write_bytes(csv_bytes)
        write_json(RAW / "ipeds" / f"{key}_dict.json", tables[table])
        record(manifest, f"ipeds_{key}", url, data, year=year, member=member, release="provisional")
        manifest["ipeds_years"][key] = year
        done.add(key)
        print(f"IPEDS {table} (provisional release): {member}, {len(tables[table]['vars'])} variables")
    return done


def fetch_ipeds(manifest):
    """Each file from its own newest published year (IPEDS releases its surveys on different schedules)."""
    forced = os.environ.get("IPEDS_YEAR")
    this_year = datetime.date.today().year
    years = [int(forced)] if forced else list(range(this_year, this_year - 5, -1))
    (RAW / "ipeds").mkdir(parents=True, exist_ok=True)
    manifest["ipeds_years"] = {}
    for key, name_for in IPEDS_FILES.items():
        year = next((y for y in years if exists(f"{IPEDS}{name_for(y)}.zip")
                     and exists(f"{IPEDS}{name_for(y)}_Dict.zip")), None)
        if year is None:
            raise SystemExit(f"no published IPEDS {name_for(this_year)} file with a dictionary (tried {years})")
        name = name_for(year)
        url = f"{IPEDS}{name}.zip"
        data = get(url)
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            member = pick_member(zf, ".csv")
            (RAW / "ipeds" / f"{key}.csv").write_bytes(zf.read(member))
        record(manifest, f"ipeds_{key}", url, data, year=year, member=member)
        durl = f"{IPEDS}{name}_Dict.zip"
        ddata = get(durl)
        with zipfile.ZipFile(io.BytesIO(ddata)) as zf:
            parsed = parse_dictionary(zf.read(pick_member(zf, ".xlsx")))
        write_json(RAW / "ipeds" / f"{key}_dict.json", parsed)
        record(manifest, f"ipeds_{key}_dict", durl, ddata, year=year)
        manifest["ipeds_years"][key] = year
        print(f"IPEDS {name}: {member}, {len(parsed['vars'])} variables in dictionary")
    if not forced:
        provisional(manifest, dict(manifest["ipeds_years"]))


def fetch_scorecard(manifest):
    """Optional: Scorecard's download host refuses GitHub's runners (403). Without it, the Scorecard-only
    columns (earnings, accreditor, SAT average, ...) stay empty. On a machine it lets through, set
    SCORECARD_ZIP_FILE to a downloaded copy or let this function download it."""
    local = os.environ.get("SCORECARD_ZIP_FILE")
    url = os.environ.get("SCORECARD_ZIP_URL")
    try:
        if local:
            data, url = open(local, "rb").read(), local
        else:
            if not url:
                try:  # a newer release than the pinned one, if the data page links it
                    page = get(SCORECARD_PAGE, tries=1).decode("utf-8", "ignore")
                    links = re.findall(r"""(https://[^"'()\s]*Most-Recent-Cohorts-Institution[^"'()\s]*\.zip)""",
                                       page)
                    url = links[0] if links else None
                except Exception as e:
                    print(f"Scorecard data page not readable ({e}); using {SCORECARD_ZIP}")
                url = url or SCORECARD_ZIP
            data = get(url)
    except Exception as e:
        print(f"WARNING: College Scorecard not downloaded ({e}); Scorecard-only columns will be empty")
        manifest["scorecard_missing"] = str(e)
        return
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        members = [n for n in zf.namelist() if n.lower().endswith(".csv") and "institution" in n.lower()]
        member = (members or [pick_member(zf, ".csv")])[0]
        (RAW / "scorecard").mkdir(parents=True, exist_ok=True)
        (RAW / "scorecard" / "institutions.csv").write_bytes(zf.read(member))
    record(manifest, "scorecard", url, data, member=member)
    print(f"Scorecard: {url} -> {member}")


def main():
    manifest = {"fetched_at": datetime.datetime.now(datetime.timezone.utc).isoformat(timespec="seconds"),
                "files": {}}
    fetch_ipeds(manifest)
    fetch_scorecard(manifest)
    write_json(OUT / "manifest.json", manifest)
    print(json.dumps({k: v["url"] for k, v in manifest["files"].items()}, indent=2))


if __name__ == "__main__":
    main()
