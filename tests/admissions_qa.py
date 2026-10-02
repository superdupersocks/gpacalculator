"""QA for the admissions data scripts (scripts/admissions/build.py and match.py) on a small fixture.

Run: python3 tests/admissions_qa.py
The fixture (tests/fixtures/admissions/) mimics the IPEDS HD/ADM/IC files, their parsed dictionaries and the
College Scorecard institution file, with deliberate problems: suppressed values, an impossible score,
more admits than applicants, an alias, a typo and two posts for one college.
"""
import csv
import io
import json
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts" / "admissions"))
import audit  # noqa: E402
import build  # noqa: E402
import cds  # noqa: E402
import cds_pages  # noqa: E402
import fetch  # noqa: E402
import match  # noqa: E402
import phase2_e_import  # noqa: E402
import phase2_s_pages  # noqa: E402

FIX = ROOT / "tests" / "fixtures" / "admissions"
fails = []


def eq(label, got, want):
    ok = got == want
    print(f"{'PASS' if ok else 'FAIL'} {label}: {got!r}" + ("" if ok else f" (want {want!r})"))
    if not ok:
        fails.append(label)


with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    build.main(["--raw", str(FIX / "raw"), "--out", str(tmp)])
    rows = {r["unitid"]: r for r in csv.DictReader(open(tmp / "institutions.csv"))}
    full_rows = rows  # with the Scorecard file; the checkpoint E checks below use these
    src = json.loads((tmp / "field_sources.json").read_text())
    report = (tmp / "qa_report.md").read_text()

    h = rows["166027"]
    eq("universe skips admin units (sector 0)", sorted(rows), ["100751", "166027", "888001", "999001", "999003"])
    eq("Harvard admit rate from IPEDS counts", h["admit_rate"], "0.0364")
    eq("Harvard yield", h["yield_rate"], "0.8377")
    eq("Harvard admissions source", (h["admissions_source"], h["admissions_year"]), ("IPEDS ADM", "2024"))
    eq("Harvard SAT ERW p25/p50/p75", (h["sat_erw_p25"], h["sat_erw_p50"], h["sat_erw_p75"]), ("740", "760", "780"))
    eq("Harvard ACT composite median", h["act_comp_p50"], "35")
    eq("Harvard test policy decoded", h["req_test_scores"], "Considered but not required")
    eq("Harvard essay requirement", h["req_essay"], "Required")
    eq("Harvard AP credit / dual credit", (h["ap_credit"], h["dual_credit"]), ("Yes", "No"))
    eq("Harvard control and level", (h["control"], h["level"]), ("Private not-for-profit", "Four or more years"))
    eq("Harvard Carnegie uses newest C??BASIC", h["carnegie"], "Doctoral Universities: Very High Research Activity")
    eq("Harvard negative longitude kept", h["lon"], "-71.118")
    eq("Harvard undergrads from IPEDS (not total enrollment)", h["undergrad_enrollment"], "7240")
    eq("Harvard tuition from the latest IC_AY year", (h["tuition_in_state"], h["tuition_out_of_state"]),
       ("59320", "59320"))
    eq("Harvard net price from the private variable", h["net_price"], "13900")
    eq("Harvard net price by income from SFA titles", h["net_price_0_30k"], "1000")
    eq("optional DRVIC absent leaves cost of attendance empty", h["cost_in_state_on_campus"], "")
    eq("Harvard grad and retention rates as 0-1", (h["grad_rate"], h["retention_rate"]), ("0.97", "0.98"))
    eq("Harvard religious affiliation not applicable -> empty", h["religious_affiliation"], "")
    eq("Harvard not-applicable dates empty", (h["closed_date"], h["merged_into"]), ("", ""))

    a = rows["100751"]
    eq("Alabama impossible SAT math 900 dropped", (a["sat_math_p25"], a["sat_math_p75"]), ("540", ""))
    eq("Alabama suppressed earnings empty, not 0", a["median_earnings_10yr"], "")
    eq("Alabama grad rate falls back to Scorecard when IPEDS is empty", a["grad_rate"], "0.72")
    eq("Alabama imputed retention dropped, Scorecard fallback", a["retention_rate"], "0.87")
    eq("Alabama public net price", a["net_price"], "21500")
    eq("each fallback measure names its source", (h["undergrad_enrollment_source"], a["grad_rate_source"],
                                                  a["retention_rate_source"], a["net_price_source"]),
       ("IPEDS", "College Scorecard", "College Scorecard", "IPEDS"))
    eq("suppressed earnings have no source column", "median_earnings_10yr_source" in a, False)

    s = rows["888001"]
    eq("Scorecard-only row uses Scorecard admissions", (s["admissions_source"], s["admit_rate"]),
       ("College Scorecard", "0.65"))
    eq("Scorecard-only row has no IPEDS counts", s["applicants"], "")
    eq("Scorecard predominant degree decoded", s["predominant_degree"], "Associate's")

    c = rows["999003"]
    eq("more admits than applicants dropped", (c["applicants"], c["admits"], c["admit_rate"]), ("", "", ""))
    eq("religious affiliation decoded", c["religious_affiliation"], "Roman Catholic")
    eq("closed college flagged", (rows["999001"]["active"], rows["999001"]["closed_date"]), ("No", "2024-06-30"))

    eq("field sources cite ADM year and variable", (src["applicants"]["source"], src["applicants"]["variable"]),
       ("IPEDS ADM2024", "APPLCN"))
    eq("ADMCON mapped by dictionary title", src["req_legacy"]["variable"], "ADMCON12")
    eq("measure sources name the variable", (src["tuition_in_state"]["variable"], src["net_price"]["variable"]),
       ("CHG2AY3", "NPGRN2 / NPIST2"))
    eq("report lists dropped values", "more admits than applicants" in report and "outside 200-800" in report, True)

    res = {r["slug"]: r for r in match.main(["--colleges", str(FIX / "colleges"),
                                             "--institutions", str(tmp / "institutions.csv"),
                                             "--out", str(tmp)])}
    eq("exact match", (res["harvard"]["method"], res["harvard"]["unitid"]), ("exact", "166027"))
    eq("'The' and 'of' ignored", (res["alabama"]["method"], res["alabama"]["unitid"]), ("exact", "100751"))
    eq("St. -> Saint", (res["saint-example"]["method"], res["saint-example"]["unitid"]), ("exact", "999003"))
    eq("alias match", (res["example-college"]["method"], res["example-college"]["unitid"]), ("alias", "999003"))
    eq("two posts for one college flagged", res["example-college"].get("duplicate_of"), "saint-example")
    eq("typo matched fuzzily", (res["closed-example"]["method"], res["closed-example"]["unitid"]),
       ("fuzzy", "999001"))
    eq("no candidate", res["gone-forever"]["method"], "none")
    review = list(csv.DictReader(open(tmp / "match_review.csv")))
    eq("review list", sorted(r["slug"] for r in review), ["example-college", "gone-forever", "saint-example"])

# Without the Scorecard file (its host refuses GitHub's runners) the build still runs on IPEDS alone.
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    shutil.copytree(FIX, tmp / "fix")
    shutil.rmtree(tmp / "fix" / "raw" / "scorecard")
    build.main(["--raw", str(tmp / "fix" / "raw"), "--out", str(tmp)])
    rows = {r["unitid"]: r for r in csv.DictReader(open(tmp / "institutions.csv"))}
    h = rows["166027"]
    eq("IPEDS-only build keeps admissions and cost", (h["admit_rate"], h["tuition_in_state"]), ("0.0364", "59320"))
    eq("IPEDS-only build leaves Scorecard columns empty", (h["median_earnings_10yr"], h["accreditor"]), ("", ""))
    eq("IPEDS-only universe", sorted(rows), ["100751", "166027", "999001", "999003"])

# Common Data Set parsing: the three-column C11 table (total column wins) and C12 split across lines.
CDS_TEXT = """C11 Percentage of all enrolled, degree-seeking, first-time, first-year (freshman) students ...
Score Included Score Not Included Total All Students
Percent who had GPA of 4.0 60.10% 70.00% 65.06%
Percent who had GPA between 3.75 and 3.99 25.00% 20.00% 22.50%
Percent who had GPA between 3.50 and 3.74 10.00% 6.00% 8.00%
Percent who had GPA between 3.25 and 3.49 3.00% 2.00% 2.50%
Percent who had GPA between 3.00 and 3.24 1.00% 1.00% 1.00%
Percent who had GPA between 2.50 and 2.99 0.90% 1.00% 0.94%
Percent who had GPA between 2.0 and 2.49 0% 0% 0%
Percent who had GPA between 1.0 and 1.99 0% 0% 0%
Percent who had GPA below 1.0 0% 0% 0%
C12 Average high school GPA of all degree-seeking, first-time, first-year (freshman) students who submitted
GPA: 4.18
Percent of total first-time, first-year (freshman) students who submitted high school GPA: 82%"""
g = cds.parse(CDS_TEXT)
eq("CDS C12 average and submit %", (g.get("gpa_avg"), g.get("gpa_submit_pct")), (4.18, 82.0))
eq("CDS C11 total column", (g.get("gpa_4_0"), g.get("gpa_250_299"), g.get("gpa_below_100")), (65.06, 0.94, 0.0))
eq("CDS bands add up", cds.check(g), [])
g = cds.parse("C11\nPercent who had GPA of 4.0 0.4\nPercent who had GPA between 3.75 and 3.99 0.6\nC12")
eq("CDS Excel fractions become percents", (cds.check(g), g["gpa_4_0"]), ([], 40.0))
eq("CDS impossible average flagged", cds.check({"gpa_avg": 39.2}), [("gpa_avg", "average GPA 39.2 outside 1-5")])
g = cds.parse("C9 Percent and number of first-time, first-year students enrolled ...\n"
              "Submitting SAT Scores 54% 1046\nSubmitting ACT Scores 21% 405\n"
              "SAT Composite 1500 1540 1570\nSAT Evidence-Based Reading and Writing 740 760 780\n"
              "SAT Math 760 780 800\nACT Composite 34 35 36\nC10")
eq("CDS C9 percentiles and submit %", (g.get("sat_comp_p50"), g.get("sat_erw_p25"), g.get("act_comp_p75"),
                                       g.get("sat_submit_pct")), (1540, 740, 36, 54.0))
eq("CDS form values decoded", (cds.clean("level", "/VI"), cds.clean("yn", "/Y"), cds.clean("pct", "0.25"),
                               cds.clean("count", "1,234"), cds.clean("gpa", "N/A")),
   ("Very Important", "Yes", 25.0, 1234, None))
form = cds.from_form({"FRSH_GPA": "3.91", "Q111_3": "/VI", "AP_RECD_1ST_N": "1000", "AP_ADMT_1ST_N": "250",
                      "AD_EDEC": "/Y", "EN_FRSH_GPA_1_P": "0.6", "EN_FRSH_GPA_2_P": "0.4"})
eq("CDS form fields (same names every template year)", {k: v for k, v in form.items() if k != "_summed"},
   {"applicants": 1000, "admits": 250, "factor_gpa": "Very Important", "gpa_avg": 3.91, "ed_offered": "Yes",
    "gpa_4_0": 60.0, "gpa_375_399": 40.0})
form = cds.from_form({"AP_RECD_1ST_N": "0", "AP_RECD_1ST_MEN_N": "5,831", "AP_RECD_1ST_WMN_N": "6,310",
                      "AP_ADMT_1ST_N": "0", "EN_TOT_1ST_N": "0", "EN_TOT_1ST_FT_MEN_N": "2890",
                      "EN_TOT_1ST_PT_WMN_N": "10"})
eq("CDS form: an uncalculated 0 total is its lines by sex added up",
   ({k: form.get(k) for k in ("applicants", "admits", "enrolled")}, sorted(form["_summed"])),
   ({"applicants": 12141, "admits": None, "enrolled": 2900}, ["applicants", "enrolled"]))
src = {"unitid": "1", "name": "X", "cds_year": "2023-24", "source_url": "https://x.edu/cds.pdf"}
vals, prov, rev = cds.decide(src, {"admissions_year": "2023", "sat_math_p25": "760", "applicants": "500"},
                             {"factor_gpa": "Important"}, {"gpa_avg": 3.9, "sat_math_p50": 780},
                             {"C.1201": "3.90", "C.911": "760", "C.912": "790", "C.117": "500", "C.1202": "88"})
eq("CDS decide: form, text+collegedata, collegedata+IPEDS",
   {k: vals.get(k) for k in ("factor_gpa", "gpa_avg", "sat_math_p25", "applicants")},
   {"factor_gpa": "Important", "gpa_avg": 3.9, "sat_math_p25": 760, "applicants": 500})
eq("CDS decide: disagreement and single readings go to review",
   sorted(r["field"] for r in rev), ["gpa_submit_pct", "sat_math_p50"])
for producer in ("tier1_xlsx", "tier2_acroform", "tier4_docling"):
    vals2, _, _ = cds.decide(src, None, {}, {}, {"C.701": "Very Important", "C.117": "2885"}, producer)
    eq(f"CDS decide: collegedata.fyi alone is not published ({producer})", vals2, {})
vals2, _, rev2 = cds.decide(src, None, {}, {"gpa_avg": 3.62, "gpa_submit_pct": 97.0, "_tight": {"gpa_avg"}}, {})
eq("CDS decide: a GPA on its label's line stands alone; a loose one doesn't",
   (vals2, [r["field"] for r in rev2]), ({"gpa_avg": 3.62}, ["gpa_submit_pct"]))
vals2, _, rev2 = cds.decide(src, {"admissions_year": "2023", "applicants": "28111"}, {"applicants": 0},
                            {"applicants": 0}, {"C.117": "0"})
eq("CDS decide: a 0 applicant count is blank, not published", (vals2, rev2), ({}, []))
vals2, _, rev2 = cds.decide(src, {"admissions_year": "2023", "applicants": "28111"}, {"applicants": 12000}, {}, {})
eq("CDS decide: C1 far from IPEDS for the same fall goes to review",
   (vals2, [r["field"] for r in rev2]), ({}, ["applicants"]))

# Which college a file belongs to: a file listed for three campuses is used for the one whose first-year class
# it matches; a single file far from IPEDS's class describes another campus; collegedata.fyi's row-shifted counts
# (fewer applicants than enrollees, or a lone count) don't disqualify a college's own file; applications can
# swing from year to year while the class size holds.
def _src(u, name, archive):
    return {"unitid": u, "name": name, "cds_year": "2025-26", "source_url": "https://x.edu/" + archive,
            "archive_url": "https://archive/" + archive + ".pdf"}


readings = {"1": (_src("1", "Main", "aaa"), {}, {"applicants": 46000}, {}, ""),
            "2": (_src("2", "Branch", "aaa"), {}, {"applicants": 46000}, {}, ""),
            "3": (_src("3", "Other", "bbb"), {}, {}, {"C.116": "52703", "C.117": "31701", "C.118": "7272"}, ""),
            "4": (_src("4", "Misread", "ccc"), {}, {"applicants": 45000}, {"C.116": "2885"}, ""),
            "5": (_src("5", "No IPEDS", "ddd"), {}, {"applicants": 900}, {}, ""),
            "6": (_src("6", "Shifted", "eee"), {}, {}, {"C.116": "1284"}, ""),
            "7": (_src("7", "Upside down", "fff"), {}, {}, {"C.116": "1829", "C.117": "900", "C.118": "2657"}, ""),
            "8": (_src("8", "Surge", "ggg"), {"applicants": 4133, "enrolled": 249}, {}, {}, ""),
            "9": (_src("9", "Two totals", "hhh"), {},
                  {"applicants": 2512, "enrolled": 162, "_alt": {"applicants": 2694, "enrolled": 364}}, {}, "")}
ipeds = {"1": {"applicants": "44000"}, "2": {"applicants": "5100"}, "3": {"applicants": "3005", "enrolled": "666"},
         "4": {"applicants": "45409"}, "6": {"applicants": "28232"}, "7": {"applicants": "9568", "enrolled": "800"},
         "8": {"applicants": "1747", "enrolled": "260"}, "9": {"applicants": "2694", "enrolled": "364"}}
skip = cds.assign(readings, ipeds)
eq("CDS files matched to their college", sorted(skip), ["2", "3"])
eq("CDS skipped file says why, and whose reading it is", skip["3"].split("(")[1].split(")")[0],
   "collegedata.fyi reading of the file: 7,272 enrolled, 52,703 applicants; IPEDS: 666 enrolled, 3,005 applicants")

LAYOUT = """C1. Applications
Total first-time, first-year students who applied in Fall 2023 1,578.0 2,055.0 9.0
                                              IN-STATE   OUT-OF-STATE   INTERNATIONAL   UNKNOWN        TOTAL
Total first-time, first-year (degree seeking) who applied     3,046      465      131                3,642
Total first-time, first-year (degree-seeking) who were admitted                                       2,933
C7. Relative importance
                                                                                          Not
                                    Very Important          Important        Considered
                                                                                       Considered
      Rigor of secondary school record      X
      Academic GPA                                 ☐                    ☐
                                                                        ✔                 ☐                 ☐
      Interview                                                                                         X
C8: SAT and ACT Policies
C21. Early decision: Does your institution offer an early decision plan for fall enrollment?
     ☐ Yes  ✔ No
Number of early decision applications received by your institution:            1,053
CDS-C Page 6
C12. Average high school GPA of all degree-seeking, first-time, first-year students who
     submitted GPA:                                                  3.71
Percent of total first-time, first-year students who submitted high school GPA:     96.4%"""
g = cds.parse(LAYOUT)
eq("CDS layout text: C1 total column, C7 marks under their headings, C21, C12",
   {k: g.get(k) for k in ("applicants", "admits", "factor_rigor", "factor_gpa", "factor_interview", "ed_offered",
                          "ed_applicants", "gpa_avg", "gpa_submit_pct")},
   {"applicants": 3642, "admits": 2933, "factor_rigor": "Very Important", "factor_gpa": "Important",
    "factor_interview": "Not Considered", "ed_offered": "No", "ed_applicants": 1053, "gpa_avg": 3.71,
    "gpa_submit_pct": 96.4})
eq("CDS layout text: a by-sex row isn't a total", cds.parse(LAYOUT.split("IN-STATE")[0]).get("applicants"), None)
C1_BLANK = """C1 Applications
Total first-time, first-year men who applied                         5,831
Total first-time, first-year women who applied                       6,310
Total first-time, first-year another gender who applied                 12
Total first-time, first-year unknown gender who applied
Total full-time, first-time, first-year men who enrolled             1,402
Total part-time, first-time, first-year men who enrolled                 3
Total full-time, first-time, first-year women who enrolled           1,488
Total first-time, first-year students who applied
Total first-time, first-year students who enrolled"""
g = cds.parse(C1_BLANK)
eq("CDS text: a blank C1 total is its lines by sex added up",
   ({k: g.get(k) for k in ("applicants", "admits", "enrolled")}, sorted(g["_summed"])),
   ({"applicants": 12153, "admits": None, "enrolled": 2893}, ["applicants", "enrolled"]))
C1_2025 = """Total first-time, first-year males who applied          14,020
Total first-time, first-year females who applied        18,734
Total first-time, first-year students of unknown sex who applied     0
First-Time, First-Year Student Applicants   In-State   Out-of-State   International   Unknown   Total
Total first-time, first-year (degree-seeking) who applied    3,512    24,960    4,282"""
eq("CDS text: a 2025-26 residency row with its Total blank isn't the total",
   cds.parse(C1_2025).get("applicants"), 32754)
eq("CDS text: a residency row whose Total adds up stays",
   cds.parse(C1_2025 + "    32,754").get("applicants"), 32754)
# Vanderbilt's workbook adds up each block of lines by sex on the row under it; a blank line by sex doesn't take
# that check sum.
_vu = cds.parse("""Total first-time, first-year males who applied                    21595
Total first-time, first-year females who applied                  26601
Total first-time, first-year students of unknown sex who applied
                                                                  48196
Total first-time, first-year (degree-seeking) who applied    2885   36130   9181   0   48196""")
eq("CDS text: a check sum under a blank line by sex isn't counted",
   (_vu.get("applicants"), _vu["_summed"], _vu["_alt"]), (48196, set(), {}))
eq("CDS text: the C1 lines as read, for the run log", _vu["_c1_lines"][-1],
   "Total first-time, first-year (degree-seeking) who applied 2885 36130 9181 0 48196")
eq("CDS text: lines \"of another gender\" count toward the total", cds.parse(
    "Total first-time, first-year men who enrolled 220\nTotal first-time, first-year women who enrolled 227\n"
    "Total first-time, first-year of another gender who enrolled 29\n"
    "Total first-time, first-year (degree-seeking) enrolled   95   347   34   476")["_alt"], {})
# Rose-Hulman puts each row's total first, in parentheses, and words its lines by sex "first-year (freshman) men".
_rh = cds.parse("""Total first-time, first-year (freshman) men who applied 4,553
Total first-time, first-year (freshman) women who applied 1,541 (Total: 6,097)
Total full-time, first-time, first-year (freshman) men who enrolled 451 (Fall Enrollment Snapshot: 451)
Total full-time, first-time, first-year (freshman) women who enrolled 153 (Fall Enrollment Snapshot: 153)
Total first-time, first-year (degree seeking) who applied (6,097) 1,295 3,847 955
Total first-time, first-year (degree seeking) enrolled (604) 159 377 68""")
eq("CDS text: the total its row adds up to, wherever it sits; \"(freshman)\" lines by sex",
   (_rh.get("applicants"), _rh.get("enrolled"), _rh["_alt"]), (6097, 604, {"applicants": 6094}))
eq("CDS text: a stated total is the number the rest of its row adds up to",
   [cds.stated_total(r) for r in ([2885, 36130, 9181, 0, 48196], [3512, 24960, 4282], [10339], [0])],
   [48196, None, 10339, None])
# A lone number in the residency row can be its in-state column (RIT); a file that states two totals keeps both,
# and a second reading settles which one is published.
_rit = cds.parse("""Total first-time, first-year men who applied 18303
Total first-time, first-year women who applied 13224
Total first-time, first-year (degree-seeking) who applied        10339""")
eq("CDS text: a file's two totals both kept", (_rit.get("applicants"), _rit["_alt"]), (10339, {"applicants": 31527}))
_src25 = {"unitid": "1", "name": "X", "cds_year": "2025-26", "source_url": "https://x.edu/cds.pdf"}
vals2, prov2, _ = cds.decide(_src25, {"admissions_year": "2024", "applicants": "27911"}, {}, _rit, {"C.116": "31527"})
eq("CDS decide: the total a second reading confirms is published",
   (vals2.get("applicants"), [p["method"] for p in prov2 if p["field"] == "applicants"]),
   (31527, ["file text + collegedata.fyi (lines by sex added up)"]))
eq("CDS text: enrollees by sex aren't counted again by full- and part-time", cds.parse(
    "Total first-time, first-year males who enrolled 100\nTotal first-time, first-year females who enrolled 120\n"
    "Total full-time, first-time, first-year males who enrolled 90\nTotal part-time, first-time, first-year males "
    "who enrolled 10\nTotal full-time, first-time, first-year females who enrolled 120").get("enrolled"), 220)
# The 2025-26 spreadsheet template indexes every question in columns right of the form, on unrelated rows: the
# C11 band rows end with the application closing month and day.
_wb = openpyxl.Workbook()
_ws = _wb.active
_ws.append(["C11", "Percentage of all enrolled ..."])
for _i, (_band, _p) in enumerate([("of 4.0", 40.0), ("between 3.75 and 3.99", 30.0), ("between 3.50 and 3.74", 20.0),
                                  ("between 3.25 and 3.49", 6.0), ("between 3.00 and 3.24", 4.0)]):
    _ws.append([None, "Percent who had GPA " + _band, None, _p, _p, _p] + [None] * 20
               + [f"C.14{_i:02d}", "Application closing date (fall): Month", _i + 1])
for _i in range(6):
    _ws.append([None] * 26 + [f"C.15{_i:02d}", "Another question", 15])
_ws.append(["C12", "Average high school GPA of all degree-seeking, first-time, first-year students who submitted GPA:",
            None, 3.71])
_buf3 = io.BytesIO()
_wb.save(_buf3)
g = cds.parse(cds.text_of(_buf3.getvalue(), "xlsx"))
eq("CDS spreadsheet: the question index right of the form is cut off",
   ([g.get(c) for c in cds.BAND_COLS[:5]], "gpa_350_374" in g["_tight"]), ([40.0, 30.0, 20.0, 6.0, 4.0], True))
_m = {c: "file text, checked" for c in cds.BAND_COLS}
_bands = dict(zip(cds.BAND_COLS, (87.6, 7.1, 3.4, 0.8, 0.7, 0.3, 0.1, 0.0, 0.0)))
eq("CDS GPA bands that can't give their average: the lone reading goes",
   cds.gpa_conflict({"gpa_avg": 3.71, **_bands}, {**_m, "gpa_avg": "form fields"}),
   [("gpa_bands", "the GPA bands put the class near 3.96, well above its 3.71 average")])
_both = {c: "file text + collegedata.fyi" for c in cds.BAND_COLS}
eq("CDS GPA bands and average both confirmed stay",
   cds.gpa_conflict({"gpa_avg": 3.71, **_bands}, {**_both, "gpa_avg": "form fields"}), [])
eq("CDS weighted average above its bands is fine",
   cds.gpa_conflict({"gpa_avg": 4.54, **dict(zip(cds.BAND_COLS, (56.8, 26.9, 9.9, 2.8, 2.9, 0.7, 0.1, 0, 0)))},
                    {**_m, "gpa_avg": "file text, checked"}), [])
eq("CDS C1 for another fall: an admit rate far from IPEDS's goes to review",
   [c for c, _ in cds.c1_jump({"applicants": 6785, "admits": 2612, "enrolled": 1742},
                              {"applicants": "8824", "admits": "7266", "enrolled": "2044"})],
   ["applicants", "admits", "enrolled"])
eq("CDS C1 for another fall: growth with the same admit rate stays",
   cds.c1_jump({"applicants": 9678, "admits": 8778, "enrolled": 519},
               {"applicants": "5037", "admits": "4492", "enrolled": "603"}), [])
# A file that contradicts itself: Appalachian State's form gives a SAT math 25th percentile of 354 beside a
# composite 25th of 1140; Pratt admitted more from its wait list than accepted a place; Barnard answers "No" to a
# wait-list policy and reports wait-list counts. The values that can't all be right go to review together.
_app = {"sat_comp_p25": 1140, "sat_comp_p50": 1200, "sat_comp_p75": 1270, "sat_erw_p25": 570, "sat_erw_p50": 610,
        "sat_erw_p75": 660, "sat_math_p25": 354, "sat_math_p50": 550, "sat_math_p75": 620}
eq("CDS consistency: SAT composite far from its sections", [c for c, _ in cds.consistency(_app)], ["sat_scores"])
eq("CDS consistency: SAT sections that add up are fine",
   cds.consistency({**_app, "sat_math_p25": 534}), [])
eq("CDS consistency: wait list and early decision contradictions", [c for c, _ in cds.consistency(
    {"waitlist_offered": 1001, "waitlist_accepted": 139, "waitlist_admitted": 925, "waitlist_policy": "No",
     "ed_offered": "No", "ed_applicants": 300, "ed_admits": 120})],
   ["waitlist_counts", "waitlist_policy", "ed_offered"])
vals2, _, rev2 = cds.decide(src, None, {**_app, "factor_gpa": "Important"}, {}, {})
eq("CDS decide: a self-contradicting group goes to review whole",
   (sorted(vals2), [r["field"] for r in rev2]), (["factor_gpa"], ["sat_scores"]))
_wb = openpyxl.Workbook()
_ws = _wb.active
_ws.append(["C11", "Percent who had GPA of 4.0", 0.55])
_ws.append([None, "Percent who had GPA between 3.75 and 3.99", 0.45])
_ws.append(["C12", "Percent of total first-time, first-year students who submitted high school GPA:", 1])
for _c in ("C1", "C2", "C3"):
    _ws[_c].number_format = "0.0%"
_buf2 = io.BytesIO()
_wb.save(_buf2)
g = cds.parse(cds.text_of(_buf2.getvalue(), "xlsx"))
eq("CDS spreadsheet percent cells read as percents", (g.get("gpa_4_0"), g.get("gpa_submit_pct")), (55.0, 100.0))

_buf = io.BytesIO()
with zipfile.ZipFile(_buf, "w") as _z:
    _z.writestr("word/document.xml", "<w:document><w:body><w:tbl><w:tr><w:tc><w:p><w:t>Average high school GPA of "
                "all degree-seeking, first-time, first-year students who submitted GPA:</w:t></w:p></w:tc><w:tc><w:p>"
                "<w:t>3.87</w:t></w:p></w:tc></w:tr></w:tbl></w:body></w:document>")
eq("CDS Word file read", cds.parse(cds.text_of(_buf.getvalue())).get("gpa_avg"), 3.87)

# IPEDS provisional release: the newest Tablesdoc on the Access page, its titles and labels per table, and
# sources marked as provisional.
page = ('<a href="/ipeds/tablefiles/tableDocs/IPEDS202324Tablesdoc.xlsx">2023-24</a>'
        '<a href="https://nces.ed.gov/ipeds/tablefiles/zipfiles/IPEDS_2023-24_Final.zip">2023-24</a>'
        '<a href="https://nces.ed.gov/ipeds/tablefiles/zipfiles/IPEDS_2024-25_Provisional.zip">2024-25</a>'
        '<a href="https://nces.ed.gov/ipeds/tablefiles/tableDocs/IPEDS202425Tablesdoc.xlsx">2024-25</a>')
eq("newest release on the Access page", fetch.newest_release(page),
   (2024, "https://nces.ed.gov/ipeds/tablefiles/tableDocs/IPEDS202425Tablesdoc.xlsx",
    "https://nces.ed.gov/ipeds/tablefiles/zipfiles/IPEDS_2024-25_Provisional.zip"))
eq("no release on the page", fetch.newest_release("<html></html>"), (None, None, None))
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Tables24"
ws.append(["TableName", "TableTitle"])
ws.append(["ADM2024", "Admissions"])
ws = wb.create_sheet("vartable24")
ws.append(["TableName", "varName", "varTitle", "DataType"])
ws.append(["ADM2024", "applcn", "Applicants total", "N"])
ws.append(["ADM2024", "ADMCON7", "Admission test scores", "N"])
ws = wb.create_sheet("valuesets24")
ws.append(["TableName", "varName", "Codevalue", "valueLabel"])
ws.append(["ADM2024", "ADMCON7", 1.0, "Required"])
buf = io.BytesIO()
wb.save(buf)
doc = fetch.parse_tablesdoc(buf.getvalue())
eq("Tablesdoc titles per table", doc["ADM2024"]["vars"]["APPLCN"]["title"], "Applicants total")
eq("Tablesdoc code labels", doc["ADM2024"]["codes"]["ADMCON7"], {"1": "Required"})
eq("provisional table without net price is skipped", fetch.usable("sfa", {"vars": {}}),
   ["average net price-students awarded grant or scholarship aid"])
eq("provisional ADM with the counts is used", fetch.usable("adm", {"vars": {v: {"title": ""} for v in
                                                                         ("APPLCN", "ADMSSN", "ENRLT")}}), [])
with tempfile.TemporaryDirectory() as tmp:
    tmp = Path(tmp)
    shutil.copytree(FIX, tmp / "fix")
    mf = tmp / "fix" / "manifest.json"
    m = json.loads(mf.read_text())
    m.setdefault("files", {})["ipeds_adm"] = {"release": "provisional"}
    m["files"]["ipeds_drvgr"] = {"release": "provisional"}
    mf.write_text(json.dumps(m))
    build.main(["--raw", str(tmp / "fix" / "raw"), "--out", str(tmp)])
    src = json.loads((tmp / "field_sources.json").read_text())
    eq("provisional ADM source marked", src["applicants"]["source"], "IPEDS ADM2024 provisional release")
    eq("provisional measure source marked", src["grad_rate"]["source"],
       "IPEDS DRVGR 2024 provisional release (else Scorecard)")
    eq("complete-file source unmarked", src["name"]["source"], "IPEDS HD2024")

# Phase 2 audit: a closure needs a closing date or a federal flag; older directories identify a post only by its
# exact name, city and state; a college still listed counts as merged only if its current record says so.


def _hd(uid, name, city, st, **kw):
    return {"UNITID": uid, "INSTNM": name, "IALIAS": "", "CITY": city, "STABBR": st, "CLOSEDAT": "-2",
            "NEWID": "-2", "DEATHYR": "-2", "CYACTIVE": "1", **kw}


_hist = {"100": {2015: _hd("100", "Old Tech Institute", "Akron", "OH"),
                 2016: _hd("100", "Old Tech Institute", "Akron", "OH", CLOSEDAT="12/31/2016", DEATHYR="2017")},
         "200": {2018: _hd("200", "Gone College", "Erie", "PA", NEWID="300")},
         "300": {2018: _hd("300", "Big University", "Erie", "PA")},
         "400": {2019: _hd("400", "Quiet School", "Boise", "ID")},
         "500": {2012: _hd("500", "Still Here College", "Mesa", "AZ", NEWID="999")},
         "810": {2018: _hd("810", "Chain Institute-North", "Kent", "OH", NEWID="820", DEATHYR="2019")},
         "820": {2018: _hd("820", "Chain Institute", "Kent", "OH")},
         "830": {2013: _hd("830", "Capital College", "Salem", "OR")},
         "850": {2017: _hd("850", "Plains Area Technical College", "Topeka", "KS")},
         "900": {2022: _hd("900", "North Valley College", "Dover", "DE", NEWID="910", DEATHYR="2023")},
         "910": {2022: _hd("910", "South Valley College", "Dover", "DE")},
         "930": {2014: _hd("930", "Old Campus", "Ames", "IA", NEWID="940", DEATHYR="2015")},
         "940": {2014: _hd("940", "State Branch", "Ames", "IA"),
                 2016: _hd("940", "State Branch", "Ames", "IA", CLOSEDAT="06/30/2016")},
         "950": {2014: _hd("950", "River Tech-Waco", "Waco", "TX", NEWID="960", DEATHYR="2015")},
         "970": {2011: _hd("970", "Hill College-Bloominton", "Bloomington", "IN")}}


def _c(uid, name, city, st="", **kw):
    return {"unitid": uid, "name": name, "city": city, "state": st, "merged_into": "", **kw}


_cur = {"300": _c("300", "Big University", "Erie"),
        "500": _c("500", "Still Here College", "Mesa", active="Yes", operating="Yes"),
        "600": _c("600", "Shut College", "Troy", closed_date="05/11/2024", active="No"),
        "700": _c("700", "Absorbed College", "Erie", merged_into="300", active="No"),
        "820": _c("820", "Chain Institute", "Kent", "OH"),
        "840": _c("840", "Capital Community College", "Salem", "OR"),
        "850": _c("850", "Plains State Tech", "Topeka", "KS"),
        "860": _c("860", "Plains Technical Institute", "Topeka", "KS"),
        "870": _c("870", "Brand College-Lakeside", "Lakeside", "NC", alias="Brand College"),
        "880": _c("880", "Brand College-Hilltop", "Hilltop", "NC"),
        "890": _c("890", "Maybe College", "Reno", "NV", active="Yes", operating="No"),
        "910": _c("910", "Valley State College", "Dover", "DE"),
        "960": _c("960", "River Tech", "Waco", "TX"),
        "98000101": _c("98000101", "Hill College-Bloomington", "Bloomington", "IN")}
_rows = [("old-tech", "Old Tech Institute", "Akron", "OH", "none", ""),
         ("gone", "Gone College", "Erie", "PA", "none", ""),
         ("quiet", "Quiet School", "Boise", "ID", "review", ""),
         ("big-u", "Big University", "Erie", "PA", "exact", "300"),
         ("still", "Still Here College", "Mesa", "AZ", "exact", "500"),
         ("still-2", "Still Here College Mesa", "Mesa", "AZ", "fuzzy", "500"),
         ("shut", "Shut College", "Troy", "NY", "exact", "600"),
         ("absorbed", "Absorbed College", "Erie", "PA", "exact", "700"),
         ("chain-north", "Chain Institute North", "Kent", "OH", "renamed", "820"),
         ("chain-south", "Chain Institute South", "Kent", "OH", "renamed", "820"),
         ("capital-college", "Capital College", "Salem", "OR", "renamed", "840"),
         ("capital-cc", "Capital Community College", "Salem", "OR", "exact", "840"),
         ("plains-atc", "Plains Area Technical College", "Topeka", "KS", "renamed", "860"),
         ("brand-college", "Brand College", "Hilltop", "NC", "alias", "870"),
         ("brand-college-lakeside", "Brand College Lakeside", "Lakeside", "NC", "exact", "870"),
         ("maybe", "Maybe College", "Reno", "NV", "exact", "890"),
         ("north-valley", "North Valley College", "Dover", "DE", "none", ""),
         ("south-valley", "South Valley College", "Dover", "DE", "review", ""),
         ("brand-college-2", "Brand College", "Hilltop", "NC", "review", ""),
         ("old-campus", "Old Campus", "Ames", "IA", "none", ""),
         ("river-tech", "River Tech", "Waco", "TX", "exact", "960"),
         ("river-tech-waco", "River Tech Waco", "Waco", "TX", "renamed", "960"),
         ("hill-bloominton", "Hill College Bloominton", "Bloomington", "IN", "fuzzy", "98000101")]
_posts = [{"slug": s, "title": t, "url": f"https://gpacalculator.net/admissions/{s}/", "city": c, "state": st,
           "fields": {"location": f"{c}, {st}", "enrollment": "100" if s == "still" else ""}}
          for s, t, c, st, _, _ in _rows]
_matches = [{"slug": s, "title": t, "location": f"{c}, {st}", "method": m, "unitid": u, "candidates": ""}
            for s, t, c, st, m, u in _rows]
_a = audit.classify(_posts, _matches, _cur, _hist)
_url = "https://gpacalculator.net/admissions/"
eq("audit: closed posts need a closing date, from an older directory or the current one",
   [(r["slug"], r["closed_on"]) for r in _a["closed"]], [("old-tech", "12/31/2016"), ("shut", "05/11/2024")])
eq("audit: merged posts go to the successor's page, a page pending review, or retire with the successor",
   [(r["slug"], r["successor_unitid"], r["treatment"].split(" once")[0]) for r in _a["merged"]],
   [("gone", "300", f"301 to {_url}big-u/"), ("absorbed", "300", f"301 to {_url}big-u/"),
    ("north-valley", "910", f"301 to {_url}south-valley/"),
    ("old-campus", "940", "retire like a closure: State Branch is no longer listed either")])
eq("audit: unmatched posts keep what IPEDS shows, left alone",
   [(r["slug"], r["finding"].split(":")[0], r["treatment"]) for r in _a["unmatched"]],
   [("quiet", "left IPEDS", "leave unchanged pending identity review"),
    ("capital-college", "left IPEDS", "leave unchanged pending identity review"),
    ("south-valley", "consolidated", "leave unchanged pending identity review"),
    ("brand-college-2", "name plus city", "leave unchanged pending identity review")])
eq("audit: a name-plus-city finding names the post already holding that IPEDS ID",
   _a["unmatched"][-1]["finding"], "name plus city: IPEDS lists UNITID 880 as Brand College-Hilltop; also the IPEDS "
   "ID of brand-college")
eq("audit: posts named for a campus record that merged into the college stay matched, as duplicates",
   [(r["slug"], r["filled_fields"]) for r in _a["duplicates"]],
   [("still", 2), ("still-2", 1), ("chain-north", 1), ("chain-south", 1), ("river-tech", 1), ("river-tech-waco", 1)])
eq("audit: a not-operating flag without a closing date stays unconfirmed",
   [r["slug"] for r in _a["unconfirmed"]], ["maybe"])
eq("audit: a campus now reported under a parent keeps its match to the campus record",
   "hill-bloominton" in {r["slug"] for k in ("closed", "merged", "unmatched", "corrections") for r in _a[k]}, False)
eq("audit: Phase 1 matches that IPEDS names contradict are reassigned",
   [(r["slug"], r["phase1_unitid"], r["unitid"], r["outcome"]) for r in _a["corrections"]],
   [("capital-college", "840", "830", "no longer listed: now in unmatched.csv"),
    ("plains-atc", "860", "850", "matched to this college instead"),
    ("brand-college", "870", "880", "matched to this college instead")])

# Checkpoint E: the fields an imported page gets, and which confident matches wait
_e = phase2_e_import
_years = {"enrollment": "2024", "net_price": "2022–23", "credits": "2024–25", "ipeds_release": "2024–25 provisional",
          "scorecard_release": "June 10, 2026"}
_h = _e.row_for({"slug": "harvard", "title": "Harvard University"}, full_rows["166027"], _years)
eq("E: Harvard's identity, place and type", (_h["ipeds_unitid"], _h["location"], _h["owning"]),
   ("166027", "Cambridge, Massachusetts", "Private nonprofit, 4-year"))
eq("E: Harvard's acceptance rate keeps a decimal under 10%, with its counts and year",
   (_h["acceptance_rate"], _h["adm_admits"], _h["adm_applicants"], _h["adm_year"]), ("3.6%", "1966", "54008", "2024"))
eq("E: Harvard's medians and requirement wording come straight from IPEDS",
   (_h["sat_reading_50"], _h["admission_requirements_high_school_gpa"]), ("760", "Considered but not required"))
eq("E: enrollment and net price carry their years; private colleges have no in-state scope",
   (_h["enrollment"], _h["enrollment_year"], _h["net_price"], _h["net_price_year"], _h["net_price_scope"]),
   ("7,240", "2024", "$13,900", "2022–23", ""))
eq("E: the average SAT is labeled College Scorecard's", (_h["average_sat_score"], _h["average_sat_score_source"]),
   ("1540", "College Scorecard"))
eq("E: fields with no fresh source are emptied", {k: _h[k] for k in _e.EMPTIED}, {k: "" for k in _e.EMPTIED})
eq("E: a public college's net price is the in-state one",
   _e.row_for({"slug": "alabama", "title": "Alabama"}, full_rows["100751"], _years)["net_price_scope"], "in-state")
_none = _e.row_for({"slug": "few", "title": "Few"}, dict(full_rows["166027"], admits="0", admit_rate="0.0"), _years)
eq("E: no acceptance rate when a college admitted none of its applicants", (_none["acceptance_rate"],
   _none["adm_admits"]), ("", "0"))
_sc = _e.row_for({"slug": "sc", "title": "Scorecard Only"}, full_rows["888001"], _years)
eq("E: College Scorecard admissions, enrollment and net price (year not stated) are left out",
   (_sc["acceptance_rate"], _sc["adm_year"], _sc["enrollment"], _sc["net_price"], _sc["average_sat_score"]),
   ("", "", "", "", "1050"))
eq("E: pct_text rounding", [_e.pct_text(v) for v in ("0.0364", "0.43", "0.0996", "0.05")],
   ["3.6%", "43%", "10%", "5%"])
_m = {"slug": "x", "unitid": "1"}
eq("E: hold reasons, in order",
   [_e.hold_reason(_m, i, c, set(), set()).split(" (")[0].split(":")[0] for i, c in (
       (full_rows["166027"], {}), (full_rows["999001"], {}), (full_rows["999003"], {}), (full_rows["888001"], {}),
       (full_rows["166027"], {"x": {"outcome": "matched to this college instead", "unitid": "2"}}), (None, {}))],
   ["", "IPEDS lists a closing date", "open in IPEDS 2024 but missing from College Scorecard's June 2026 release",
    "no IPEDS 2024 record of its own", "the audit corrected this match", "UNITID 1 isn't in institutions.csv"])

_branch = {"16884704": {"name": "Baker College of Cadillac", "opeid": "00229504", "control": ""},
           "168847": {"name": "Baker College", "opeid": "00229500", "control": "Private not-for-profit"},
           "501211": {"name": "Ohio Business College-Columbus", "opeid": "02158507", "control": ""},
           "203720": {"name": "Ohio Business College-Sheffield", "opeid": "02158500", "control": "Private for-profit"}}
_why = "no IPEDS 2024 record of its own: only College Scorecard lists it"
eq("E: a branch campus points to the college that reports for it, by UNITID or main-campus OPEID",
   [_e.next_step(_why, u, _branch, {"203720": "ohio-business-college-sheffield"}) for u in ("16884704", "501211")],
   ["leave unchanged until Baker College (UNITID 168847) has a page here, then 301 to it",
    "301 to Ohio Business College-Sheffield's page, /admissions/ohio-business-college-sheffield/, which has the "
    "federal figures for this campus"])

# After E: pages for the colleges that held pages point to, and the redirects to them
_s = phase2_s_pages
_inst = {"168847": {"name": "Baker College", "closed_date": "", "operating": "Yes", "control": "Private not-for-profit",
                    "opeid": "00229500"},
         "16884704": {"name": "Baker College of Cadillac", "closed_date": "", "operating": "", "control": "",
                      "opeid": "00229504"},
         "484613": {"name": "University of Phoenix-Arizona", "closed_date": "", "operating": "Yes",
                    "control": "Private for-profit", "opeid": "02088800"},
         "484631": {"name": "University of Phoenix-California", "closed_date": "", "operating": "No",
                    "control": "Private for-profit", "opeid": "02088801"},
         "484710": {"name": "University of Phoenix-Nevada", "closed_date": "06/05/2023", "operating": "",
                    "control": "Private for-profit", "opeid": "02088802"},
         "133997": {"name": "Florida Career College-Miami", "closed_date": "", "operating": "",
                    "control": "Private for-profit", "opeid": "02186200"},
         "203720": {"name": "Ohio Business College-Sheffield", "closed_date": "", "operating": "Yes",
                    "control": "Private for-profit", "opeid": "02158500"},
         "501211": {"name": "Ohio Business College-Columbus", "closed_date": "", "operating": "", "control": "",
                    "opeid": "02158507"}}
eq("S: a merged college's page goes to its successor's new page, Phoenix's small state units to the one Phoenix page; "
   "a closed or non-operating successor retires it; Fortis keeps its page under its new IPEDS ID",
   [_s.c_action({"slug": slug, "successor_unitid": u}, _inst, {}) for slug, u in (
       ("baker-college-of-owosso", "168847"), ("university-of-phoenix-san-diego-campus", "484631"),
       ("university-of-phoenix-las-vegas-campus", "484710"), ("florida-career-college-clearwater", "133997"),
       ("fortis-institute", "494436"))],
   [("301", "168847"), ("301", "484613"),
    ("retire", "merged into University of Phoenix-Nevada, which closed 06/05/2023 (IPEDS)"),
    ("retire", "merged into Florida Career College-Miami, which College Scorecard (June 2026) doesn't list as "
               "operating"),
    ("rematch", "494436")])
eq("S/M: a branch campus goes to the college that reports for it, when that college has or gets a page",
   [_s.e_action({"why": w, "unitid": u}, _inst, {"203720": "ohio-business-college-sheffield"}) for w, u in (
       (_why, "16884704"), (_why, "501211"), ("open in IPEDS 2024 but missing from College Scorecard", "203720"))],
   ["168847", "203720", ""])
_pages, _s.PAGES = _s.PAGES, {"166027": ("harvard-new", "Harvard New")}
_rows = dict(full_rows)
_rows.update({"16602701": dict(full_rows["166027"], control="", opeid="00215501"), "100751": full_rows["100751"]})
_new, _srows, _act = _s.build([], [{"slug": "harvard-extension", "why": _why, "unitid": "16602701"}], _rows, {},
                              _years)
_s.PAGES = _pages
eq("S: each new page gets E's fields, and its campuses a 301 to it",
   (_new, [(r["slug"], r["post_title"], r["ipeds_unitid"], r["acceptance_rate"]) for r in _srows],
    [(a["checkpoint"], a["slug"], a["action"], a["target"]) for a in _act]),
   ([{"slug": "harvard-new", "post_title": "Harvard New"}], [("harvard-new", "Harvard New", "166027", "3.6%")],
    [("S", "harvard-extension", "301", "https://gpacalculator.net/admissions/harvard-new/")]))

# CDS files on another college's website aren't this college's (same-named colleges)
_web = {"219718": {"name": "Bethel University", "website": "www.bethelu.edu/"},
        "173160": {"name": "Bethel University", "website": "https://www.bethel.edu"},
        "190512": {"name": "CUNY Bernard M Baruch College", "website": "www.baruch.cuny.edu"},
        "190600": {"name": "CUNY Hunter College", "website": "www.hunter.cuny.edu"}}
_bs = cds.sites(_web)
eq("CDS: a file on another college's website is not used; its own site, a subdomain or Drive is fine",
   [cds.elsewhere({"unitid": u, "source_url": url}, _web, _bs) for u, url in (
       ("219718", "https://www.bethel.edu/ir/cds-2425.xlsx"), ("173160", "https://www.bethel.edu/ir/cds-2425.xlsx"),
       ("219718", "https://ir.bethelu.edu/cds.pdf"), ("190512", "https://drive.usercontent.google.com/x"),
       ("190512", "https://www.hunter.cuny.edu/cds.pdf"))],
   [["Bethel University"], [], [], [], ["CUNY Hunter College"]])

# A CDS file off the college's website counts once a page on that website links it
_drive = {"unitid": "1", "source_url": "https://drive.usercontent.google.com/download?id=1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF&export=download",
          "archive_url": "https://x/sources/a/2025-26/" + "ab" * 32 + ".pdf", "cds_year": "2024-25"}
_cdn = {"unitid": "2", "source_url": "https://live-csu-northridge.pantheonsite.io/sites/default/files/2025-02/CDS%202024.xlsx",
        "archive_url": "", "cds_year": "2024-25"}
_box = {"unitid": "3", "source_url": "https://public.boxcloud.com/d/1/b1!abc", "archive_url": "", "cds_year": "2025-26"}
eq("CDS pages: what names the file in a page",
   [(t["google"], t["paths"], t["sha256"][:4]) for t in map(cds_pages.tokens, (_drive, _cdn, _box))],
   [("1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF", [], "abab"),
    ("", ["2025-02/CDS 2024.xlsx", "2025-02/CDS%202024.xlsx"], ""), ("", [], "")])
_page = ('<p>Common Data Set: <a href="https://drive.google.com/file/d/1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF/view">'
         '2024-2025</a> <a href="/sites/default/files/2025-02/CDS%202024.xlsx">CDS 2024-25 (Excel)</a>'
         ' <a href="https://upenn.box.com/s/k3j2h1">Common Data Set 2025-2026</a> <a href="/ir/">IR</a></p>')
_pl = cds_pages.links(_page, "https://www.csun.edu/ir/")
eq("CDS pages: a Drive ID or the CDN path in the page finds its link; Box needs the year and a download",
   [cds_pages.found_in(_page, _pl, cds_pages.tokens(s)) for s in (_drive, _cdn, _box)]
   + [cds_pages.cds_links(_pl, "2025-26"), cds_pages.download_url("https://upenn.box.com/s/k3j2h1")],
   [("https://drive.google.com/file/d/1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF/view", "Google file ID"),
    ("https://www.csun.edu/sites/default/files/2025-02/CDS%202024.xlsx", "file path"), None,
    ["https://upenn.box.com/s/k3j2h1"], "https://upenn.box.com/shared/static/k3j2h1"])
eq("CDS pages: only files off the college's own site (and not commondataset.org) need a page",
   [cds_pages.needs_page({"source_url": u}, w) for u, w in (
       ("https://www.csun.edu/x.pdf", "www.csun.edu"), ("https://ir.csun.edu/x.pdf", "https://www.csun.edu/"),
       ("https://drive.usercontent.google.com/download?id=x", "www.csun.edu"),
       ("https://commondataset.org/x", "www.csun.edu"))],
   [False, False, True, False])

_site = {"https://www.example.edu/": '<a href="/about/">About</a> <a href="/offices/ir/">Institutional Research</a>',
         "https://www.example.edu/offices/ir/": '<a href="cds/">Common Data Set</a> <a href="/news/">News</a>',
         "https://www.example.edu/offices/ir/cds/": '<a href="https://drive.google.com/file/d/'
                                                    '1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF/view">CDS 2024-2025</a>',
         "https://www.example.edu/about/": "<p>About us</p>"}
_crawl = cds_pages.Crawler(10)
_crawl.robots = {h: __import__("urllib.robotparser").robotparser.RobotFileParser() for h in ("www.example.edu",)}
_crawl.robots["www.example.edu"].parse([])
_site["https://www.example.edu/"] = '<a href="/ir/IR Overview.pptx">Institutional research overview</a> ' \
    + _site["https://www.example.edu/"]
_asked = []
_crawl.get = lambda url, check=True, limit=0: _asked.append(url) or _site.get(url, "").encode() or None
_row, _read = _crawl.find({**_drive, "name": "Example College"}, "www.example.edu")
eq("CDS pages: the crawl follows institutional-research and CDS links to the page that links the file",
   (_row["page_url"], _row["match"], _read), ("https://www.example.edu/offices/ir/cds/", "Google file ID", 3))
eq("CDS pages: the crawl doesn't open documents as pages",
   [u for u in _asked if "pptx" in u.lower()], [])
eq("CDS pages: addresses with spaces or accents are sent percent-encoded, escapes kept",
   [cds_pages.fetchable(u) for u in ("https://www.x.edu/facet/docs/FACET Effort Overview.pptx",
                                     "https://www.x.edu/a%20b/caf\u00e9/ ", "https://www.x.edu/p?a=1&b=c d")],
   ["https://www.x.edu/facet/docs/FACET%20Effort%20Overview.pptx", "https://www.x.edu/a%20b/caf%C3%A9/",
    "https://www.x.edu/p?a=1&b=c%20d"])


# A page counts where it ends up after redirects: on the college's site, cited there, its links read from there
_moved = cds_pages.Crawler(10)
_moved.robots = {h: _crawl.robots["www.example.edu"] for h in ("www.example.edu", "www.elsewhere.org")}
_moved_to = {"https://www.example.edu/": "https://www.example.edu/home/",
             "https://www.example.edu/ir/": "https://www.elsewhere.org/ir/"}
_moved_site = {"https://www.example.edu/": '<a href="ir-office/">Institutional Research</a> <a href="/ir/">IR</a>',
               "https://www.example.edu/home/ir-office/": '<a href="https://drive.google.com/file/d/'
                                                         '1cb-7QPm2EL_CSJP4lw1RN1qLEKfHiQiF/view">CDS 2024-25</a>',
               "https://www.example.edu/ir/": _site["https://www.example.edu/offices/ir/cds/"]}


def _moved_get(url, check=True, limit=0):
    _moved.final[url] = _moved_to.get(url, url)
    return _moved_site.get(url, "").encode() or None


_moved.get = _moved_get
_row2, _ = _moved.find({**_drive, "name": "Example College"}, "www.example.edu")
_moved_site["https://www.example.edu/"] = '<a href="/ir/">IR</a>'
_row3, _ = _moved.find({**_drive, "name": "Example College"}, "www.example.edu")
eq("CDS pages: a redirected page is read and cited at its final address, and only if that is on the college's site",
   (_row2["page_url"], _row3), ("https://www.example.edu/home/ir-office/", None))
eq("CDS pages: the SHA-256 to match is the copy cds.py read, else the archive copy's name",
   [cds_pages.tokens(s)["sha256"][:4] for s in ({**_drive, "read_sha256": "cd" * 32}, _drive)], ["cdcd", "abab"])
eq("CDS: a Google Sheets export (new bytes on every download) is read from the archived copy first",
   [[k for k, _ in cds.copies(s)] for s in (
       {"source_url": "https://doc-00-60-sheets.googleusercontent.com/export/x/y/1/2/*/z?format=xlsx",
        "archive_url": "https://a/b"},
       {"source_url": "https://drive.usercontent.google.com/download?id=x", "archive_url": "https://a/b"},
       {"source_url": "https://www.x.edu/cds.pdf", "archive_url": ""})],
   [["archive_url", "source_url"], ["source_url", "archive_url"], ["source_url"]])

def _refuse(*a, **k):
    raise __import__("http.client").client.InvalidURL("URL can't contain control characters")


_urlopen, cds_pages.urllib.request.urlopen = cds_pages.urllib.request.urlopen, _refuse
try:
    _got = cds_pages.Crawler(5).get("https://www.example.edu/x", check=False)
finally:
    cds_pages.urllib.request.urlopen = _urlopen
_boom = cds_pages.Crawler(5)
_boom.find = lambda src, website: 1 / 0
eq("CDS pages: an address that can't be read is skipped, and an error in one college's search ends only that one",
   (_got, cds_pages.find_safely(_boom, _drive, "www.example.edu")),
   (None, (None, 0, "ZeroDivisionError: division by zero")))

print("\nALL PASSED" if not fails else f"\nFAILED: {len(fails)}")
sys.exit(1 if fails else 0)
